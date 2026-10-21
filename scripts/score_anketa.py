"""Оценка анкет по рубрике (правила + LLM) и проверка качества на исторических данных.

Запуск:
    python scripts/score_anketa.py --provider none          # только правила, без LLM
    python scripts/score_anketa.py --provider auto --runs 3 # GigaChat, запасной OpenRouter (ключи в .env)
    python scripts/score_anketa.py --provider auto --limit 10   # пробный прогон на 10 анкетах

Результат: data/scored.xlsx — листы «Баллы», «Проверка», «Рубрика».
Ответы LLM кэшируются в data/llm_cache: повторный прогон тех же анкет бесплатен.
"""
import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from scipy.stats import spearmanr  # noqa: E402
from sklearn.metrics import roc_auc_score  # noqa: E402

from scorer.extract import assess  # noqa: E402
from scorer.llm import LLMError, make_provider  # noqa: E402
from scorer.score import aggregate, candidate_from_clean, item_scores, load_rubric  # noqa: E402


def score_all(df, rubric, provider, runs, temperature):
    rows = []
    for n, (_, r) in enumerate(df.iterrows(), 1):
        c = candidate_from_clean(r.to_dict())
        llm = assess(c, rubric, provider, runs, temperature) if provider else None
        s = item_scores(c, rubric, llm)
        agg = aggregate(s, rubric)
        row = dict(id=c["id"], cohort=r["cohort"], screen_score=r["screen_score"], stage=r["stage"],
                   taken=r["taken"], hours_per_year=r["hours_per_year"])
        row.update({f"s_{k}": v for k, v in s.items()})
        row.update({f"F_{k}": v for k, v in agg["factors"].items()})
        row.update(index=agg["index"], coverage=agg["coverage"], grade=agg["grade"])
        if llm and llm.get("ok"):
            row.update(exp_categories="; ".join(f"{k}: «{q}»" for k, q in llm["categories"].items()),
                       **{f"evidence_{k}": v for k, v in llm["evidence"].items()},
                       summary=llm["summary"], recommended_direction=llm["recommended_direction"],
                       interview_questions="\n".join(llm["interview_questions"]),
                       red_flags="; ".join(llm["red_flags"]))
        row["flags"] = "; ".join(llm["flags"]) if llm else ""
        rows.append(row)
        if provider:
            print(f"  {n}/{len(df)} {c['id']}: индекс {agg['index']}, оценка {agg['grade']}", flush=True)
    return pd.DataFrame(rows)


def auc(y, x):
    d = pd.DataFrame({"y": y, "x": x}).dropna()
    if d.y.nunique() < 2 or d.x.nunique() < 2:
        return None, len(d)
    return round(roc_auc_score(d.y.astype(int), d.x.astype(float)), 3), len(d)


def rho(a, b):
    d = pd.DataFrame({"a": a, "b": b}).dropna()
    if d.a.nunique() < 2 or d.b.nunique() < 2:
        return None, len(d)
    return round(spearmanr(d.a, d.b).statistic, 3), len(d)


def evaluate(s, rubric):
    """Связь индекса и отдельных пунктов с решениями кураторов и с реальной активностью."""
    taken = s[s.taken == 1]
    active = (taken.hours_per_year >= 20).astype(float).where(taken.hours_per_year.notna())
    f25 = s[s.cohort == "2025-весна"]
    rows = []

    def add(name, value_n, what):
        rows.append(dict(Метрика=name, Значение=value_n[0], n=value_n[1], Что_значит=what))

    add("Spearman: индекс ↔ балл куратора", rho(s["index"], s.screen_score), "насколько рубрика повторяет кураторов")
    add("AUC: индекс → приглашение (весна 2025)", auc(f25.screen_score > 0, f25["index"]),
        "в листах 2024 года нет непрошедших, поэтому только 2025")
    add("AUC: индекс → активность ≥ 20 ч/год (взятые)", auc(active, taken["index"]), "главная метрика: реальный успех")
    add("AUC: балл куратора → активность ≥ 20 ч/год (взятые)", auc(active, taken.screen_score), "базовая линия, которую надо обогнать")
    for item in rubric["items"]:
        col = f"s_{item['id']}"
        if s[col].notna().any():
            add(f"пункт {item['id']} (вес {item['weight']}): Spearman с баллом куратора", rho(s[col], s.screen_score), "")
            add(f"пункт {item['id']}: AUC → активность ≥ 20 ч", auc(active, taken[col]), "")
    for f in rubric["factors"]:
        add(f"фактор {f}: AUC → активность ≥ 20 ч", auc(active, taken[f"F_{f}"]), "")
    return pd.DataFrame(rows)


def suggest_thresholds(s):
    """Пороги индекса, при которых доля оценок 0/1/2/3 совпадает с долей у кураторов."""
    d = s[["index", "screen_score"]].dropna()
    if len(d) < 20:
        return None
    shares = d.screen_score.value_counts(normalize=True).sort_index().cumsum()
    return [round(float(d["index"].quantile(shares.get(g, 0))), 1) for g in (0, 1, 2)]


def write_xlsx(path, sheets):
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E9EBEF")
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for j, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 100)), 1):
            w = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[get_column_letter(j)].width = min(max(8, w + 2), 50)
    wb.save(path)


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", default=str(ROOT / "data" / "anketa_clean.csv"))
    ap.add_argument("--out", default=str(ROOT / "data" / "scored.xlsx"))
    ap.add_argument("--rubric", default=str(ROOT / "config" / "rubric.yaml"))
    ap.add_argument("--provider", default="none", choices=["none", "auto", "gigachat", "openrouter"])
    ap.add_argument("--runs", type=int, default=3)
    ap.add_argument("--temperature", type=float, default=0.3)
    ap.add_argument("--limit", type=int)
    a = ap.parse_args()

    rubric = load_rubric(a.rubric)
    df = pd.read_csv(a.input)
    if a.limit:
        df = df.head(a.limit)
    provider = None
    if a.provider != "none":
        try:
            provider = make_provider(a.provider)
        except LLMError as e:
            sys.exit(str(e))
        print(f"Провайдер: {provider.name} ({provider.model}), прогонов на анкету: {a.runs}")

    s = score_all(df, rubric, provider, a.runs, a.temperature)
    ev = evaluate(s, rubric)
    th = suggest_thresholds(s) if provider else None
    rub = pd.DataFrame([dict(id=i["id"], фактор=rubric["factors"][i["factor"]], вес=i["weight"], метод=i["method"],
                             предложено_командой="да" if i.get("proposed") else "", вопрос=i["question"])
                        for i in rubric["items"]])
    write_xlsx(a.out, {"Баллы": s, "Проверка": ev, "Рубрика": rub})

    print(f"\nАнкет: {len(s)}; средний coverage: {s.coverage.mean():.2f}")
    print(ev.head(4).to_string(index=False))
    if th:
        print(f"\nПредложенные пороги оценки 0–3 (как у кураторов): {th}; сейчас в рубрике {rubric['grade_thresholds']}")
    if provider:
        print(f"LLM: новых вызовов {provider.calls}, из кэша {provider.hits}")
    print(f"Файл: {a.out}")


if __name__ == "__main__":
    main()
