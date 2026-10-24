"""Сведение и очистка таблиц музея с прошлыми наборами волонтёров.

Листы исходных таблиц (список — в config/history.yaml, файлы — в data/history/) сводятся в одну таблицу
«одна строка — один кандидат», значения приводятся к единому виду, каждое смысловое исправление пишется в журнал.

Запуск (обычно через scripts/import_history.py, который заодно оценивает анкеты и переобучает модель):
    python scripts/clean_anketa.py

Результат в папке --out:
    anketa_clean.xlsx  — листы «Анкеты», «Исправления», «Вопросы представителю», «Словарь», «Исходные данные»
    anketa_clean.csv   — лист «Анкеты» в UTF-8 для pandas
"""
import argparse
import datetime as dt
import re
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent

# Дата, на которую в таблицах актуальны «Еще состоит» и «Активность» (на неё же посчитаны формулы TODAY()
# в листах 2024 года). Задаётся в config/history.yaml (ref_date), main() её подставляет.
REF_DATE = dt.date.today()


def load_history_config(path):
    """config/history.yaml → (дата актуальности таблиц, список листов)."""
    import yaml
    config = yaml.safe_load(Path(path).read_text(encoding="utf-8"))
    tables = []
    for table in config["tables"]:
        tables.append(dict(table, interview_format=table.get("interview_format") or None,
                           start_assumed=bool(table.get("start_assumed"))))
    return config.get("ref_date") or dt.date.today(), tables

# Столбцы исходных листов ищем по началу заголовка: в листах они стоят в разном порядке.
FIELDS = [
    ("ЯВКА", "attended_raw"), ("Сколько баллов", "screen_score_raw"), ("Взяли ли", "taken_raw"),
    ("Длительность", "years_raw"), ("Еще состоит", "still_raw"), ("Активность", "hours_raw"),
    ("ФИО", "fio_raw"), ("Дата рождения", "birth_raw"), ("ВУЗ", "vuz_raw"), ("Факультет", "faculty_raw"),
    ("Курс", "course_raw"), ("Кафедра", "department_raw"), ("Сфера научных", "sci_raw"),
    ("Опыт работы", "experience_raw"), ("Владение", "languages_raw"), ("Интересующие", "directions_raw"),
    ("Чего Вы ожидаете", "expectations_raw"), ("Ваше хобби", "hobby_raw"), ("В чем Ваша", "superpower_raw"),
    ("Откуда", "source_raw"),
]

MONTHS = {"январ": 1, "феврал": 2, "март": 3, "апрел": 4, "мая": 5, "май": 5, "июн": 6, "июл": 7,
          "август": 8, "сентябр": 9, "октябр": 10, "ноябр": 11, "декабр": 12}
ORDINALS = {"первый": 1, "второй": 2, "третий": 3, "четвертый": 4, "четвёртый": 4, "пятый": 5}


# ---------------------------------------------------------------- журнал исправлений

class Log:
    def __init__(self):
        self.rows = []

    def fix(self, rec, field, raw, new, rule):
        if ws(raw) == ws(new):
            return
        self.rows.append(dict(id=rec["id"], src_sheet=rec["src_sheet"], src_row=rec["src_row"],
                              field=field, raw=ws(raw), new=ws(new), rule=rule))


def ws(v):
    """Строка без лишних пробелов, переносов и неразрывных пробелов."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        v = v.strftime("%d.%m.%Y")
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def eq(a, b):
    """Сравнение, безопасное для pd.NA."""
    return not pd.isna(a) and a == b


def is_placeholder(s):
    return re.fullmatch(r"[\s\-—–.,_]*", s) is not None


def text(rec, log, field, raw, extra_placeholders=()):
    s = ws(raw)
    if is_placeholder(s) or s.lower() in extra_placeholders:
        log.fix(rec, field, raw, "", "прочерк или пустое значение → пусто")
        return ""
    return s


# ---------------------------------------------------------------- исходы

def parse_outcomes(rec, raw, cfg, log, flags, notes):
    score = raw["screen_score_raw"]
    rec["screen_score"] = int(score) if isinstance(score, (int, float)) else pd.NA

    # Явка: столбец есть только в листах 2024 года.
    att = ws(raw.get("attended_raw"))
    if "attended_raw" not in raw:
        rec["attended"] = pd.NA
    elif att.lower().startswith("да"):
        rec["attended"] = 1
        extra = re.sub(r"^да+\s*", "", att, flags=re.I).strip()
        if extra:
            notes.append(extra)
            log.fix(rec, "attended", att, "1", "пометка из ячейки «Явка» перенесена в notes")
        elif att != "да":
            log.fix(rec, "attended", att, "1", "вариант написания «да»")
    else:
        rec["attended"] = 0

    taken = ws(raw["taken_raw"]).upper()
    if taken == "ДА":
        rec["taken"] = 1
    elif taken == "—":
        rec["taken"] = 0
    else:
        rec["taken"] = 0 if eq(rec["screen_score"], 0) else pd.NA
        if rec["taken"] is pd.NA:
            flags.append("нет решения по кандидату")
            log.fix(rec, "taken", raw["taken_raw"] if raw["taken_raw"] is not None else "(пусто)", "нет данных",
                    "пустая ячейка «Взяли ли» у приглашённого кандидата")

    if eq(rec["taken"], 1) and eq(rec["attended"], 0):
        rec["attended"] = 1
        log.fix(rec, "attended", "(пусто)", "1", "явка не отмечена, но кандидата взяли → пришёл")
    if pd.isna(rec["attended"]) and eq(rec["taken"], 1):
        rec["attended"] = 1

    if eq(rec["screen_score"], 0):
        rec["stage"] = "не приглашён"
    elif pd.isna(rec["taken"]):
        rec["stage"] = "нет данных"
    elif rec["taken"] == 1:
        rec["stage"] = "взят"
    elif eq(rec["attended"], 0):
        rec["stage"] = "не пришёл"
    else:
        rec["stage"] = "не взят после собеседования"

    # Пребывание в программе.
    still = ws(raw["still_raw"]).upper()
    rec["still_in_program"] = 1 if still == "ДА" else 0 if still == "НЕТ" else pd.NA
    if still == "НЕТ" and ws(raw["still_raw"]) != "НЕТ":
        log.fix(rec, "still_in_program", raw["still_raw"], "0", "лишние пробелы в «НЕТ»")

    rec["years_in_program"] = pd.NA
    rec["exit_date"] = pd.NA
    formula = raw.get("years_formula")
    yv = raw["years_raw"]
    if formula:
        m = re.search(r'YEARFRAC\("(\d\d\.\d\d\.\d{4})",\s*(TODAY\(\)|"(\d\d\.\d\d\.\d{4})")', formula)
        start = dt.datetime.strptime(m.group(1), "%d.%m.%Y").date()
        if m.group(3):
            end = dt.datetime.strptime(m.group(3), "%d.%m.%Y").date()
            rec["exit_date"] = end
            if pd.isna(rec["still_in_program"]):
                rec["still_in_program"] = 0
                log.fix(rec, "still_in_program", "—", "0", "в формуле длительности есть дата выхода → вышел")
        else:
            end = REF_DATE
        rec["years_in_program"] = round((end - start).days / 365.25, 2)
        log.fix(rec, "years_in_program", f"{formula} = {ws(yv)}", rec["years_in_program"],
                "длительность пересчитана из формулы с точностью до 0,01 года")
    elif isinstance(yv, dt.datetime):
        flags.append("длительность испорчена")
        if eq(rec["still_in_program"], 1):
            rec["years_in_program"] = round((REF_DATE - cfg["start"]).days / 365.25, 2)
            flags.append("длительность восстановлена от условной даты старта")
            log.fix(rec, "years_in_program", f"дата {yv:%d.%m}", rec["years_in_program"],
                    f"вместо числа дата; волонтёр в программе → от {cfg['start']:%d.%m.%Y} до {REF_DATE:%d.%m.%Y}")
        else:
            log.fix(rec, "years_in_program", f"дата {yv:%d.%m}", "нет данных", "вместо числа дата")
    else:
        s = ws(yv).replace(",", ".")
        if re.fullmatch(r"\d+(\.\d+)?", s):
            rec["years_in_program"] = float(s)
            if isinstance(yv, str):
                log.fix(rec, "years_in_program", f"«{yv}» (текст)", s, "число, записанное текстом")

    if not pd.isna(rec["years_in_program"]) and cfg["start_assumed"]:
        max_years = (REF_DATE - dt.date(cfg["start"].year, 1, 1)).days / 365.25
        if rec["years_in_program"] > max_years:
            flags.append("длительность больше срока с начала набора")

    # В столбце «Активность в часах в год» по словам представителя стоит СУММА часов за всё время в программе.
    hv = raw["hours_raw"]
    s = ws(hv)
    if isinstance(hv, (int, float)):
        rec["hours_total"] = float(hv)
    elif re.search(r"\d", s):
        rec["hours_total"] = float(re.search(r"\d+(?:[.,]\d+)?", s).group().replace(",", "."))
        log.fix(rec, "hours_total", s, rec["hours_total"], "число извлечено из текста")
    else:
        rec["hours_total"] = pd.NA

    # Часы в год = сумма / годы в программе. Меньше полугода не делим, чтобы не раздувать новичков.
    total, years = rec["hours_total"], rec["years_in_program"]
    if pd.isna(total) or pd.isna(years):
        rec["hours_per_year"] = pd.NA
    else:
        rec["hours_per_year"] = round(total / max(years, 0.5), 1)

    h = rec["hours_per_year"]
    if pd.isna(h):
        rec["hours_band"] = ""
    elif h < 10:
        rec["hours_band"] = "мало"
    elif h <= 30:
        rec["hours_band"] = "средний"
    else:
        rec["hours_band"] = "отлично"

    if eq(rec["taken"], 1) and pd.isna(h):
        flags.append("взят, но нет часов активности")
    if eq(rec["taken"], 0) and not pd.isna(h):
        flags.append("не взят, но есть часы активности")


# ---------------------------------------------------------------- возраст

def parse_birth(rec, raw, cfg, log, flags):
    v = raw["birth_raw"]
    s = ws(v)
    date, year = None, None
    if isinstance(v, (int, float)):
        year = int(v)
    elif isinstance(v, dt.datetime):
        date = v.date()
    else:
        m = re.fullmatch(r"(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})", s)
        if m:
            d, mo, y = map(int, m.groups())
            if y < 100:
                y += 2000 if y <= REF_DATE.year % 100 else 1900
            try:
                date = dt.date(y, mo, d)
            except ValueError:
                flags.append("дата рождения не распознана")
        else:
            m = re.fullmatch(r"(\d{1,2})\s+([а-яё]+)\s+(\d{4})", s.lower())
            mo = m and next((n for k, n in MONTHS.items() if m.group(2).startswith(k)), None)
            if mo:
                date = dt.date(int(m.group(3)), mo, int(m.group(1)))
    ref = cfg["interview_date"]
    if date:
        rec["birth_date"], rec["birth_year"], rec["birth_precision"] = date, date.year, "день"
        rec["age"] = ref.year - date.year - ((ref.month, ref.day) < (date.month, date.day))
        if date.strftime("%d.%m.%Y") != s:
            log.fix(rec, "birth_date", s, date.strftime("%d.%m.%Y"), "дата приведена к формату ДД.ММ.ГГГГ")
    elif year:
        rec["birth_date"], rec["birth_year"], rec["birth_precision"] = pd.NA, year, "год"
        rec["age"] = ref.year - year
    else:
        rec["birth_date"], rec["birth_year"], rec["birth_precision"], rec["age"] = pd.NA, pd.NA, "", pd.NA
        flags.append("нет даты рождения")
    if pd.isna(rec["age"]):
        rec["is_adult"] = pd.NA
    else:
        rec["is_adult"] = int(rec["age"] >= 18)
        if rec["birth_precision"] == "год" and rec["age"] == 18:
            flags.append("возраст 17 или 18: известен только год рождения")
        if rec["age"] < 18:
            flags.append("младше 18")


# ---------------------------------------------------------------- учёба

def parse_course(rec, raw, cfg, log, flags):
    s = ws(raw["course_raw"])
    t = s.lower()
    out = dict(status="не указано", edu_level="не указано", course_num=pd.NA, studying=pd.NA, graduation_year=pd.NA)
    m = re.search(r"(\d)\s*курс\w*\s*(бак|маг|спец)", t)
    if is_placeholder(s):
        flags.append("курс не указан")
    elif m:
        level = {"бак": "бакалавриат", "маг": "магистратура", "спец": "специалитет"}[m.group(2)]
        out.update(status="студент", edu_level=level, course_num=int(m.group(1)), studying=1)
    elif "спо" in t and "курс" in t:
        num = next((n for w, n in ORDINALS.items() if w in t), pd.NA)
        out.update(status="студент", edu_level="СПО", course_num=num, studying=1)
    elif re.search(r"дпо|дополнительного образования|учусь в центре", t) and not re.search(r"законч|оконч", t):
        out.update(status="слушатель ДПО", edu_level="ДПО / переподготовка", studying=1)
    elif "пенсионер" in t:
        out.update(status="пенсионер", studying=0)
    elif re.fullmatch(r"работаю", t):
        out.update(status="работает", studying=0)
    elif re.fullmatch(r"специалитет", t):
        out.update(edu_level="специалитет")
        flags.append("курс не указан")
    elif re.search(r"законч|оконч|выпуск|заверш|освоил|дипломирован|бакалавр|магистр", t):
        out.update(status="выпускник", edu_level="высшее (окончено)", studying=0)
        years = [int(y) for y in re.findall(r"(?:19|20)\d\d", t)]
        if years:
            out["graduation_year"] = max(years)
        elif "в этом году" in t:
            out["graduation_year"] = cfg["interview_date"].year
        if t in ("бакалавр", "магистр"):
            flags.append("неясно: учится или окончил")
    else:
        flags.append("курс не распознан")
    rec.update(out)
    rec["first_year_bach"] = int(out["edu_level"] == "бакалавриат" and eq(out["course_num"], 1))
    canon = (f"{out['course_num']} курс, {out['edu_level']}" if not pd.isna(out["course_num"])
             else out["edu_level"])
    log.fix(rec, "course", s, canon if s and not m else s, "свободный ответ о курсе разобран на поля")


# ---------------------------------------------------------------- вуз и факультет

VUZ_RULES = [
    (r"цпи", "ФГАУ ЦПИ Минпросвещения"),
    (r"ломоносов|^мгу\b|^мгу$", "МГУ"),
    (r"рггу|российский государственный гуманитарный", "РГГУ"),
    (r"вшэ|высшая школа экономики", "НИУ ВШЭ"),
    (r"просвещени|^гуп\b|мгоу|московский государственный областной", "ГУП (бывш. МГОУ)"),
    (r"гаугн|академический университет гуманитарных", "ГАУГН"),
    (r"косыгин", "РГУ им. Косыгина"),
    (r"^мпгу$", "МПГУ"),
    (r"^мгпу$", "МГПУ"),
    (r"рудн", "РУДН"),
    (r"сеченов", "Сеченовский университет"),
    (r"бауман|мвту", "МГТУ им. Баумана"),
    (r"спбгу|санкт-петербургский государственный университет", "СПбГУ"),
    (r"государственный университет управления", "ГУУ"),
    (r"финансовый университет", "Финансовый университет"),
    (r"институт культуры", "МГИК"),
    (r"социально-гуманитарный", "ГСГУ"),
    (r"колледж предпринимательства|кп№ ?11", "Колледж предпринимательства № 11"),
    (r"красноярский государственный педагогический", "КГПУ им. Астафьева"),
    (r"каподистри", "Афинский университет им. Каподистрии"),
    (r"станкин", "МГТУ «Станкин»"),
    (r"мтуси", "МТУСИ"),
    (r"институт арт", "Институт арт-бизнеса и антиквариата"),
    (r"добролюбов", "НГЛУ им. Добролюбова"),
    (r"^кфу$", "КФУ"),
]


def norm_vuz(part):
    p = part.strip(" .").lower()
    for pat, name in VUZ_RULES:
        if re.search(pat, p):
            return name
    return part.strip(" .")


def parse_vuz(rec, raw, log):
    s = text(rec, log, "vuz", raw["vuz_raw"])
    parts = [p for p in re.split(r",|;|(?:^|\s)\d\.\s*", s) if p.strip(" .")]
    names = list(dict.fromkeys(norm_vuz(p) for p in parts))
    rec["vuz"] = "; ".join(names)
    rec["vuz_multi"] = int(len(names) > 1)
    log.fix(rec, "vuz", s, rec["vuz"], "название вуза приведено к единому виду")


FIELD_GROUPS = [
    # «Искусство» раньше «истории»: иначе «Факультет истории искусств» попадёт в историю.
    (r"искусств|изящн|изобразит", "искусство и культура"),
    (r"истори|истфак|фипип|архив|археолог|антрополог|этнолог|религи|унцир|музеолог", "история"),
    (r"филолог|филфак|лингвист|язык|переводч|фияр", "филология и языки"),
    (r"философ", "философия"),
    (r"педагог|учител|дефектолог|дошкольн|методическ|образовани", "педагогика"),
    (r"журналист|сми\b|коммуникац", "медиа и журналистика"),
    (r"туризм|экскурсовод|гид", "туризм и экскурсии"),
    # Управление культурой (в т. ч. «культурная политика») музей относит к непрофильному, как арт-менеджмент.
    (r"эконом|менеджмент|управлени|сервис|культурн\w* политик", "экономика и управление"),
    (r"политолог|мировой полит|социальн|социолог", "социально-политические науки"),
    (r"юрид|право", "право"),
    (r"культур", "искусство и культура"),
    (r"медицин|провизор", "медицина"),
    (r"техни|инженер|автомат|математ|электросвяз|криоген", "технические и естественные"),
    (r"гуманитарн", "гуманитарные (прочие)"),
]
PROFILE = {"история": "профильное", "искусство и культура": "профильное",
           "филология и языки": "смежное", "философия": "смежное", "педагогика": "смежное",
           "медиа и журналистика": "смежное", "туризм и экскурсии": "смежное",
           "гуманитарные (прочие)": "смежное", "социально-политические науки": "смежное"}


def field_group(faculty, department):
    """Область по факультету; кафедра уточняет, если факультет общий («гуманитарных наук») или не распознан."""
    def match(t):
        return next((name for pat, name in FIELD_GROUPS if re.search(pat, t.lower())), None) if t else None
    g = match(faculty)
    if g is None or "гуманитарн" in faculty.lower():
        g = match(department) or g
    if g:
        return g
    return "не указано" if not (faculty or department) else "другое"


def parse_faculty(rec, raw, log):
    s = text(rec, log, "faculty", raw["faculty_raw"])
    f = s.replace("Исьтории", "Истории").replace("политологи и", "политологии и")
    low = f.lower()
    if re.search(r"фипип|истори\w*\s*,?\s*политолог\w*\s*,?\s*и права|история, политология и право", low):
        f = "Факультет истории, политологии и права"
    elif re.fullmatch(r"(исторический|истфак|история)( факультет)?\.?", low):
        f = "Исторический факультет"
    elif re.fullmatch(r"факультет гуманитарных наук", low):
        f = "Факультет гуманитарных наук"
    elif f:
        f = f[0].upper() + f[1:]
    rec["faculty"] = f
    log.fix(rec, "faculty", s, f, "написание факультета приведено к единому виду")

    d = text(rec, log, "department", raw["department_raw"], ("без кафедры",))
    d = d.replace("кибер ика", "кибернетика").replace(" ,", ",").rstrip(" .")
    if d:
        d = d[0].upper() + d[1:]
    rec["department"] = d
    rec["field_group"] = field_group(f, d)
    rec["profile_match"] = PROFILE.get(rec["field_group"], "не указано" if rec["field_group"] == "не указано"
                                       else "непрофильное")


# ---------------------------------------------------------------- языки, направления, источник

LANGS = [("древнегреч", "Древнегреческий"), ("древнеисланд", "Древнеисландский"), ("англ", "Английский"),
         ("нем", "Немецкий"), ("франц", "Французский"), ("испан", "Испанский"), ("итал", "Итальянский"),
         ("кит", "Китайский"), ("корей", "Корейский"), ("греч", "Греческий"), ("латын", "Латынь"),
         ("норвеж", "Норвежский"), ("татар", "Татарский")]


def parse_languages(rec, raw, log):
    s = text(rec, log, "languages", raw["languages_raw"], ("не владею", "только русский"))
    found, notes = [], []
    for seg in re.split(r",|\bи\b", s):
        low = seg.lower()
        if re.search(r"\b[аa][12]\b", low) or "со словар" in low:
            notes.append(f"исключено (ниже B1): {seg.strip()}")
            continue
        # Найденный ключ вырезаем, чтобы «греч» не сработал внутри «древнегреческий».
        for key, name in LANGS:
            if key in low:
                found.append(name)
                low = low.replace(key, " ")
    if re.search(r"не подтвержд|не было практики", s.lower()):
        notes.append("уровень не подтверждён")
    langs = sorted(set(found), key=lambda x: (x != "Английский", x))
    rec["languages"] = "; ".join(langs)
    rec["n_languages"] = len(langs)
    rec["lang_note"] = "; ".join(notes)
    log.fix(rec, "languages", s, rec["languages"], "список языков приведён к единому виду")


def parse_directions(rec, raw, log):
    s = ws(raw["directions_raw"])
    low = s.lower()
    rec["dir_org"] = int("организац" in low)
    rec["dir_project"] = int("проект" in low)
    rec["dir_fond"] = int("фонд" in low)
    rec["dir_inclusive"] = int("инклюз" in low)
    rec["dir_media"] = pd.NA  # в анкетах 2024–2025 гг. такого варианта не было
    rec["n_dirs"] = rec["dir_org"] + rec["dir_project"] + rec["dir_fond"] + rec["dir_inclusive"]
    rec["only_fond"] = int(rec["dir_fond"] == 1 and rec["n_dirs"] == 1)
    rec["directions_note"] = "; ".join(re.findall(r"\(([^)]*)\)", s))


def parse_source(rec, raw, log):
    s = text(rec, log, "source", raw["source_raw"])
    low = s.lower()
    if not s:
        src = "не указано"
    elif "социальн" in low:
        src = "Социальные сети"
    elif "друз" in low:
        src = "Рекомендации друзей"
    elif "сайт" in low:
        src = "Сайт ГИМ"
    elif re.search(r"препод|педагог|вуз|универ|администрац|практик", low):
        src = "Вуз или преподаватель"
    elif re.search(r"искал|поиск|интернет", low):
        src = "Сам нашёл в интернете"
    else:
        src = "Другое"
    rec["source"] = src
    rec["source_detail"] = "" if s in ("Социальные сети", "Рекомендации друзей", "Сайт Исторического музея") else s
    log.fix(rec, "source", s, src, "источник отнесён к категории")


# ---------------------------------------------------------------- сборка

def read_sheet(path, cfg):
    wf = openpyxl.load_workbook(path)[cfg["sheet"]]
    wv = openpyxl.load_workbook(path, data_only=True)[cfg["sheet"]]
    header = [ws(c.value) for c in wv[1]]
    cols = {}
    for j, h in enumerate(header, 1):
        for prefix, name in FIELDS:
            if h.startswith(prefix) and name not in cols:
                cols[name] = j
    rows = []
    for i in range(2, wv.max_row + 1):
        raw = {name: wv.cell(i, j).value for name, j in cols.items()}
        if all(ws(v) == "" for v in raw.values()):
            continue
        f = wf.cell(i, cols["years_raw"]).value
        raw["years_formula"] = f if isinstance(f, str) and f.startswith("=") else None
        rows.append((i, raw, header, [wv.cell(i, j).value for j in range(1, len(header) + 1)]))
    return rows


def load_manual_fixes(path):
    """Ручные правки ячеек из data/manual_fixes.yaml (в git не попадает: там тексты из анкет).

    Формат — список записей:
        - id: O24-013
          field: experience
          cut_from: "начало фрагмента, который нужно вырезать до конца ячейки"
          reason: "заметка куратора в ответе кандидата"
    """
    if not Path(path).exists():
        return []
    import yaml
    return yaml.safe_load(Path(path).read_text(encoding="utf-8")) or []


def apply_manual_fixes(rec, fixes, log, notes):
    """Вырезает фрагмент из ответа кандидата. Если это заметка куратора — переносит её в notes:
    модель её больше не увидит, а куратор — увидит."""
    for fix in fixes:
        if fix["id"] != rec["id"]:
            continue
        value = rec.get(fix["field"], "")
        position = value.find(fix["cut_from"])
        if position < 0:
            print(f"Ручная правка для {fix['id']}.{fix['field']} не применена: фрагмент не найден")
            continue
        removed = value[position:].strip()
        rec[fix["field"]] = value[:position].strip()
        notes.append(removed)
        log.fix(rec, fix["field"], value, rec[fix["field"]], f"ручная правка: {fix.get('reason', 'фрагмент удалён')}")


def build(folder, tables, fixes=()):
    """folder — папка с таблицами (data/history), tables — листы из config/history.yaml."""
    log = Log()
    records, raw_rows = [], []
    for cfg in tables:
        path = Path(folder) / cfg["file"]
        if not path.exists():
            raise SystemExit(f"Нет файла {path}: положите таблицу в {folder} или поправьте config/history.yaml")
        for n, (row, raw, header, values) in enumerate(read_sheet(path, cfg), 1):
            rec = dict(id=f"{cfg['prefix']}-{n:03d}", cohort=cfg["cohort"], src_sheet=cfg["sheet"], src_row=row,
                       interview_format=cfg["interview_format"] or "")
            flags, notes = [], []
            rec["fio"] = ws(raw.get("fio_raw"))
            parse_outcomes(rec, raw, cfg, log, flags, notes)
            parse_birth(rec, raw, cfg, log, flags)
            parse_course(rec, raw, cfg, log, flags)
            parse_vuz(rec, raw, log)
            parse_faculty(rec, raw, log)
            rec["sci_interests"] = text(rec, log, "sci_interests", raw["sci_raw"])
            exp = text(rec, log, "experience", raw["experience_raw"])
            if re.fullmatch(r"(нет|не имеется|не имею|без опыта|отсутствует|к сожалению,? ?пока|у опыта)\.?", exp.lower()):
                log.fix(rec, "experience", exp, "", "явный ответ «нет опыта» → пусто, статус «нет (указал)»")
                rec["experience"], rec["experience_status"] = "", "нет (указал)"
            else:
                rec["experience"] = exp
                rec["experience_status"] = "описан" if exp else "не заполнено"
            parse_languages(rec, raw, log)
            parse_directions(rec, raw, log)
            rec["expectations"] = text(rec, log, "expectations", raw["expectations_raw"])
            rec["hobby"] = text(rec, log, "hobby", raw["hobby_raw"])
            rec["superpower"] = text(rec, log, "superpower", raw["superpower_raw"])
            parse_source(rec, raw, log)
            rec["mentions_practice"] = int(bool(re.search(r"практик", f"{rec['expectations']} {rec['source_detail']}".lower())))
            apply_manual_fixes(rec, fixes, log, notes)
            rec["dq_flags"] = "; ".join(flags)
            rec["notes"] = "; ".join(notes)
            records.append(rec)
            raw_rows.append(dict(id=rec["id"], src_sheet=cfg["sheet"], src_row=row,
                                 **{h.split(" (")[0][:60] or f"col{j}": ws(v) for j, (h, v) in enumerate(zip(header, values), 1)},
                                 years_formula=raw["years_formula"] or ""))
    df = pd.DataFrame(records)[COLUMNS]
    return df, pd.DataFrame(log.rows), pd.DataFrame(raw_rows)


COLUMNS = [
    "id", "cohort", "src_sheet", "src_row", "interview_format", "fio",
    "screen_score", "attended", "taken", "stage", "years_in_program", "exit_date", "still_in_program",
    "hours_total", "hours_per_year", "hours_band",
    "birth_date", "birth_year", "birth_precision", "age", "is_adult",
    "status", "edu_level", "course_num", "first_year_bach", "studying", "graduation_year",
    "vuz", "vuz_multi", "faculty", "department", "field_group", "profile_match", "sci_interests",
    "experience", "experience_status", "languages", "n_languages", "lang_note",
    "dir_org", "dir_project", "dir_fond", "dir_inclusive", "dir_media", "n_dirs", "only_fond", "directions_note",
    "expectations", "hobby", "superpower", "source", "source_detail", "mentions_practice",
    "dq_flags", "notes",
]

DICTIONARY = [
    ("id", "Идентификатор: O24 — очное собеседование 30.10.2024, D24 — дистанционное 29.10.2024, F25 — форма весны 2025"),
    ("cohort", "Набор: 2024-осень или 2025-весна"),
    ("src_sheet, src_row", "Лист и строка исходной таблицы — по ним можно найти исходную запись"),
    ("interview_format", "Формат собеседования; для 2025 года в таблице не указан"),
    ("fio", "ФИО — заполнено только в листе «Дистант 29.10»"),
    ("screen_score", "Балл куратора после анкеты: 0 — не позвали, 1 — слабая, 2 — средняя, 3 — хорошая"),
    ("attended", "Пришёл на собеседование: 1/0; в 2025 году столбца «Явка» нет, известна только у взятых"),
    ("taken", "Взяли в команду: 1/0; пусто — в таблице нет решения"),
    ("stage", "Итог: не приглашён / не пришёл / не взят после собеседования / взят / нет данных"),
    ("years_in_program", "Лет в программе. 2024 — пересчитано из формул до 0,01 года; для действующих — на 21.10.2026"),
    ("exit_date", "Дата выхода из программы (из формул листов 2024 года)"),
    ("still_in_program", "Ещё в программе: 1/0; пусто — не применимо (не взят) или нет данных"),
    ("hours_total", "Активность: сумма часов за всё время в программе (так заполнен исходный столбец)"),
    ("hours_per_year", "Часов в год = hours_total / years_in_program; срок меньше 0,5 года считается как 0,5"),
    ("hours_band", "Шкала музея по часам в год: мало (< 10), средний (10–30), отлично (> 30)"),
    ("birth_date, birth_year, birth_precision", "Дата и год рождения; точность «день» или «год» (в 2025 году указан только год)"),
    ("age, is_adult", "Возраст на дату собеседования (для 2025 года — на 01.03.2025, ±1 год) и признак 18+"),
    ("status", "студент / выпускник / работает / пенсионер / слушатель ДПО"),
    ("edu_level, course_num", "Уровень образования и номер курса, разобранные из свободного ответа «Курс»"),
    ("first_year_bach", "1 — первый курс бакалавриата («жёлтый флаг» музея)"),
    ("studying, graduation_year", "Учится сейчас: 1/0; год окончания, если указан"),
    ("vuz, vuz_multi", "Вуз в едином написании (несколько — через «;») и признак нескольких вузов"),
    ("faculty, department", "Факультет и кафедра в едином написании"),
    ("field_group", "Область образования по факультету и кафедре"),
    ("profile_match", "Профильность для музея: профильное (история, искусство) / смежное / непрофильное"),
    ("experience_status", "описан / нет (указал явно) / не заполнено"),
    ("languages, n_languages, lang_note", "Языки от B1 (ниже B1 исключены, причина в lang_note)"),
    ("dir_*", "Выбранные направления: организационное, проектное, фондовое, инклюзивное; медиа в 2024–2025 гг. не было (пусто)"),
    ("n_dirs, only_fond", "Число выбранных направлений; 1 — выбраны только фонды"),
    ("directions_note", "Уточнения кандидата в скобках, например «квесты», «археология»"),
    ("source, source_detail", "Источник в категории и исходная формулировка для «Другое» и нестандартных ответов"),
    ("mentions_practice", "1 — кандидат упоминает практику в ожиданиях или источнике"),
    ("dq_flags", "Проблемы качества данных по строке"),
    ("notes", "Пометки кураторов из исходных ячеек"),
]


def questions(df):
    f = df["dq_flags"]
    q = []
    n = int(((df.cohort == "2025-весна") & (df.stage == "нет данных")).sum())
    if n:
        q.append(("Решение по кандидату", f"Весна 2025: у {n} приглашённых кандидатов (балл 1–3) пустая ячейка «Взяли ли». "
                  "Это «не пришли», «не взяли» или данных нет?"))
    n = int(f.str.contains("длительность испорчена").sum())
    if n:
        q.append(("Длительность", f"Весна 2025: в {n} ячейках «Длительность» вместо числа стоит дата 02.05. Все эти люди ещё "
                  "в программе. Какое число имелось в виду? Пока длительность посчитана от условной даты старта 01.04.2025."))
    q.append(("Дата старта", "Когда начался весенний набор 2025 года: дата первого дня в программе? Для осени 2024 это 01.11.2024."))
    n = int(f.str.contains("длительность больше срока").sum())
    if n:
        q.append(("Длительность", f"Весна 2025: у {n} действующего волонтёра указано 2 года в программе — больше, чем прошло с весны 2025. "
                  "Он был в программе раньше?"))
    n = int(f.str.contains("взят, но нет часов").sum())
    if n:
        q.append(("Часы активности", f"У {n} взятого волонтёра не указаны часы активности. Можно их восстановить?"))
    q.append(("Анкеты осени 2024", "В таблице осени 2024 только приглашённые на собеседование (нет баллов 0). "
              "Сохранились ли анкеты тех, кого не позвали? Без них модель приглашения учится только на весне 2025."))
    q.append(("Выход из программы", "Осень 2024: у взятых с датой выхода в формуле ячейка «Ещё состоит» пустая. "
              "Мы считаем, что они вышли. Верно?"))
    return pd.DataFrame(q, columns=["Тема", "Вопрос"])


def write_xlsx(path, sheets):
    with pd.ExcelWriter(path, engine="openpyxl", date_format="DD.MM.YYYY") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
    wb = openpyxl.load_workbook(path)
    head_fill = PatternFill("solid", fgColor="E9EBEF")
    for wsx in wb.worksheets:
        for c in wsx[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(vertical="top", wrap_text=True)
        wsx.freeze_panes = "B2"
        wsx.auto_filter.ref = wsx.dimensions
        for j, col in enumerate(wsx.iter_cols(min_row=1, max_row=min(wsx.max_row, 200)), 1):
            width = max(len(ws(c.value)) for c in col)
            wsx.column_dimensions[get_column_letter(j)].width = min(max(8, width + 2), 60)
        for row in wsx.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, (dt.date, dt.datetime)):
                    c.number_format = "DD.MM.YYYY"
    wb.save(path)


def main():
    global REF_DATE
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", default=str(ROOT / "config" / "history.yaml"))
    ap.add_argument("--folder", default=str(ROOT / "data" / "history"), help="где лежат таблицы музея")
    ap.add_argument("--out", default=str(ROOT / "data"))
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    REF_DATE, tables = load_history_config(a.config)
    fixes = load_manual_fixes(ROOT / "data" / "manual_fixes.yaml")
    df, log, raw = build(a.folder, tables, fixes)
    df.to_csv(out / "anketa_clean.csv", index=False, encoding="utf-8-sig")
    write_xlsx(out / "anketa_clean.xlsx", {
        "Анкеты": df, "Исправления": log, "Вопросы представителю": questions(df),
        "Словарь": pd.DataFrame(DICTIONARY, columns=["Столбец", "Описание"]), "Исходные данные": raw,
    })
    print(f"Кандидатов: {len(df)}; исправлений: {len(log)}; строк с флагами: {(df.dq_flags != '').sum()}")
    print(f"Файлы: {out / 'anketa_clean.xlsx'}, {out / 'anketa_clean.csv'}")


if __name__ == "__main__":
    main()
