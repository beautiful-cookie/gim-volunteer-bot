"""Excel-выгрузка набора: «Рейтинг», «Ответы», «Вопросы к собеседованию», «Сводка набора».

Файл собирается в памяти и отдаётся кнопкой «Скачать» в дашборде.
"""
from collections import Counter
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def safe_cell(value):
    """Защита от подстановки формул: ответ кандидата «=HYPERLINK(…)» Excel выполнил бы как формулу.
    Такой текст начинаем с апострофа — Excel покажет его как обычный текст."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def ranking_sheet(df, factor_titles, item_titles, decisions):
    rows = []
    for place, (_, c) in enumerate(df.sort_values("index", ascending=False).iterrows(), 1):
        decision, note = decisions.get(c["key"], ("", ""))
        row = {"Место": place, "Анкета": c["label"], "ФИО": c["name"], "Страница ВК": c["vk_url"],
               "Источник данных": c["source"], "Индекс": c["index"], "Оценка 0–3": c["grade"],
               "P(успеха)": c["p_success"]}
        for factor, title in factor_titles.items():
            row[title] = c.get(f"F_{factor}")
        for item, question in item_titles.items():
            row[f"{item}: {question[:40]}"] = c.get(f"s_{item}")
        row.update({"Флаги": "; ".join(c["flags"]), "Направления": ", ".join(c["directions"]),
                    "Рекомендуемое направление": c["recommended_direction"],
                    "Решение куратора": decision, "Комментарий": note})
        rows.append(row)
    return pd.DataFrame(rows)


def answers_sheet(df):
    rows = []
    for _, c in df.iterrows():
        row = {"Анкета": c["label"], "ФИО": c["name"]}
        for question, answer in c["answers"].items():
            if question in ("followups", "_clarified"):
                continue
            row[question] = ", ".join(answer) if isinstance(answer, list) else answer
        for number, pair in enumerate(c["answers"].get("followups", []), 1):
            row[f"Уточнение {number}"] = f"{pair[0]} — {pair[1] or 'пропущено'}"
        rows.append(row)
    return pd.DataFrame(rows)


def interview_sheet(df):
    rows = []
    for _, c in df.sort_values("index", ascending=False).iterrows():
        for number, question in enumerate(c["interview_questions"], 1):
            rows.append({"Анкета": c["label"], "ФИО": c["name"], "№": number, "Вопрос": question})
    return pd.DataFrame(rows, columns=["Анкета", "ФИО", "№", "Вопрос"])


def summary_sheet(df, decisions):
    rows = [("Анкет", len(df)), ("Средний индекс", round(df["index"].mean(), 1))]
    for grade in (3, 2, 1, 0):
        rows.append((f"Оценка {grade}", int((df["grade"] == grade).sum())))
    for direction, count in Counter(d for ds in df["directions"] for d in ds).most_common():
        rows.append((f"Направление: {direction}", count))
    for channel, count in Counter(df["channel"].fillna("")).most_common():
        rows.append((f"Источник: {channel or 'не указан'}", count))
    for flag, count in Counter(f for fs in df["flags"] for f in fs).most_common(10):
        rows.append((f"Флаг: {flag}", count))
    for decision, count in Counter(decisions.get(k, ("не решено", ""))[0] for k in df["key"]).most_common():
        rows.append((f"Решение: {decision}", count))
    return pd.DataFrame(rows, columns=["Показатель", "Значение"])


def build_excel(df, factor_titles, item_titles, decisions):
    sheets = {
        "Рейтинг": ranking_sheet(df, factor_titles, item_titles, decisions),
        "Ответы": answers_sheet(df),
        "Вопросы к собеседованию": interview_sheet(df),
        "Сводка набора": summary_sheet(df, decisions),
    }
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, table in sheets.items():
            table.map(safe_cell).to_excel(writer, sheet_name=name, index=False)

    # Оформление: жирная шапка, закреплённая первая строка, фильтр, разумная ширина столбцов.
    buffer.seek(0)
    workbook = load_workbook(buffer)
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E9EBEF")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for index, column in enumerate(sheet.iter_cols(max_row=min(sheet.max_row, 60)), 1):
            width = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[get_column_letter(index)].width = min(max(8, width + 2), 60)
    result = BytesIO()
    workbook.save(result)
    return result.getvalue()
